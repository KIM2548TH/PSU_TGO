import io
import openpyxl
from webapp.models.material_mapping_excel_model import MaterialMappingExcel
from webapp.models.materail_model import Material


def export_material_mapping_excel(
    campus_id, department_key, year, sheet_name, input_excel_path, output_excel_path
):
    # ดึง mapping doc
    mapping_doc = MaterialMappingExcel.objects(
        campus_id=campus_id,
        department_key=department_key,
        year=year,
        sheet_name=sheet_name,
    ).first()
    if not mapping_doc:
        raise Exception("ไม่พบข้อมูล Mapping")

    # โหลดไฟล์ Excel จาก DB หรือ path
    from webapp.models.file_model import TemplateExcel

    template = (
        TemplateExcel.objects(year=year, campus_id=campus_id)
        .order_by("display_name")
        .first()
    )
    if template and template.file and template.file.data:
        # โหลดจาก DB
        file_stream = io.BytesIO(template.file.data)
        wb = openpyxl.load_workbook(file_stream)
    else:
        # Fallback: ใช้ไฟล์ที่ระบุ
        wb = openpyxl.load_workbook(input_excel_path)

    ws = wb[sheet_name]

    # วนแต่ละ key ที่ mapping ไว้
    for key, selected_list in mapping_doc.mappings.items():
        total_amount = 0
        for sel in selected_list:
            form_id = sel if isinstance(sel, str) else sel.get("id")
            materials = Material.objects(
                form_and_formula=form_id,
                campus=campus_id,
                department=department_key,
                year=year,
            )
            for material in materials:
                try:
                    total_amount += float(material.result or 0)
                except Exception:
                    pass
        # หา row ที่คอลัมน์ B ตรงกับ key
        for row in ws.iter_rows(min_row=2, max_col=4):
            if str(row[1].value).strip() == key:
                row[3].value = total_amount  # ใส่ในคอลัมน์ D
                break

    wb.save(output_excel_path)
    return output_excel_path


def export_material_mapping_excel_to_download(
    campus_id, department_key, year, sheet_name, input_excel_path
):
    mapping_doc = MaterialMappingExcel.objects(
        campus_id=campus_id,
        department_key=department_key,
        year=year,
        sheet_name=sheet_name,
    ).first()
    if not mapping_doc:
        raise Exception("ไม่พบข้อมูล Mapping")

    # โหลดไฟล์ Excel จาก DB หรือ path
    from webapp.models.file_model import TemplateExcel

    template = (
        TemplateExcel.objects(year=year, campus_id=campus_id)
        .order_by("display_name")
        .first()
    )
    if template and template.file and template.file.data:
        # โหลดจาก DB
        file_stream = io.BytesIO(template.file.data)
        wb = openpyxl.load_workbook(file_stream)
    else:
        # Fallback: ใช้ไฟล์ที่ระบุ
        wb = openpyxl.load_workbook(input_excel_path)

    ws = wb[sheet_name]

    for key, selected_list in mapping_doc.mappings.items():
        total_amount = 0
        for sel in selected_list:
            form_id = sel if isinstance(sel, str) else sel.get("id")
            materials = Material.objects(
                form_and_formula=form_id,
                campus=campus_id,
                department=department_key,
                year=year,
            )
            for material in materials:
                try:
                    total_amount += float(material.result or 0)
                except Exception:
                    pass
        for row in ws.iter_rows(min_row=2, max_col=4):
            if str(row[1].value).strip() == key:
                row[3].value = total_amount
                break

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
